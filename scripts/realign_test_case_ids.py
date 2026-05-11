"""Realign sheet IDs and titles in C1SE.14-Test-Case-Sprint1.xlsx to match
the Sprint Backlog (C2SE.76_Product&SprintBacklog.xlsx) and the Test Plan
section 5.1 (C2SE.76_TestPlan_Ver1.1.docx).

Only the following are changed:
  - sheet name (workbook tab)
  - sheet metadata row 1 col C (sheet title) and row 2 col C (sheet code)
  - TC ID column (column B) for each test-case row

Test content (Type, Description, Pre-Condition, Step, Data, Expected,
Actual Result, Status, Note) is preserved 1:1.

Mapping derives from:
  • Test Plan §5.1 / Table 14: FC10 Certification Enrollment & Progress,
    FC11 Practice Engine, FC14 Exam Simulation, FC17 Repository & Content
    Import, FC19 Evidence and Reporting for Certification.
  • Sprint Backlog: TEA-11..TEA-17 (Sprint 2), STU-12..STU-21 (Sprint 3),
    ADM-09..ADM-14 (Sprint 5).
"""

from openpyxl import load_workbook

WB_PATH = r"e:\UpLoad\capstoneprojectedvision\Ed_Vision\C1SE.14-Test-Case-Sprint1.xlsx"

# (old_sheet_name, new_sheet_name, new_sheet_code, new_tc_prefix)
RENAMES = [
    ("TEACHER-01 Exam Import",           "S2-FC17-TEA12 Exam Import",        "S2-FC17-TEA12", "TC-S2F17-EX"),
    ("TEACHER-02 Diagnostic Import",     "S2-FC17-TEA16 Diagnostic Import",  "S2-FC17-TEA16", "TC-S2F17-DG"),
    ("TEACHER-03 Practice Import",       "S2-FC17-TEA15 Practice Import",    "S2-FC17-TEA15", "TC-S2F17-PR"),
    ("TEACHER-04 Vocabulary Import",     "S2-FC17-TEA15 Vocab Import",       "S2-FC17-TEA15", "TC-S2F17-VC"),
    ("STUDENT-01 Intake Test",           "S3-FC10-STU21 TOEIC Intake",       "S3-FC10-STU21", "TC-S3F10-TI"),
    ("STUDENT-02 Learning Map",          "S3-FC11-STU12 Learning Map",       "S3-FC11-STU12", "TC-S3F11-LM"),
    ("STUDENT-03 Mock Exam",             "S3-FC14-STU15 Mock Exam",          "S3-FC14-STU15", "TC-S3F14-MX"),
    ("STUDENT-04 Leaderboard & Streak",  "S5-FC19-ADM09 Leaderboard",        "S5-FC19-ADM09", "TC-S5F19-LB"),
    ("STUDENT-05 IELTS Intake",          "S3-FC10-STU21 IELTS Intake",       "S3-FC10-STU21", "TC-S3F10-II"),
    ("STUDENT-06 IELTS Practice",        "S3-FC11-STU12 IELTS Practice",     "S3-FC11-STU12", "TC-S3F11-IP"),
    ("STUDENT-07 IELTS Band Test",       "S3-FC14-STU15 IELTS Band Test",    "S3-FC14-STU15", "TC-S3F14-BT"),
]

DATA_START_ROW = 12  # rows 12..N hold test cases per the template


def main() -> None:
    wb = load_workbook(WB_PATH)
    summary = []

    for old_name, new_name, sheet_code, tc_prefix in RENAMES:
        if old_name not in wb.sheetnames:
            print(f"⚠  Skip (not found): {old_name}")
            continue

        ws = wb[old_name]

        # 1) Update header metadata
        #    Row 1 column C → sheet title, Row 2 column C → sheet code
        ws.cell(row=1, column=3, value=new_name)
        ws.cell(row=2, column=3, value=sheet_code)

        # 2) Renumber TC IDs in column B starting at DATA_START_ROW
        renumbered = 0
        for row in range(DATA_START_ROW, ws.max_row + 1):
            tc_cell = ws.cell(row=row, column=2)
            no_cell = ws.cell(row=row, column=1)  # No. column drives presence
            # Only touch rows that are actual test-case rows (have either
            # an existing TC id or a numeric No.)
            if tc_cell.value or isinstance(no_cell.value, (int, float)):
                renumbered += 1
                tc_cell.value = f"{tc_prefix}-{renumbered:02d}"

        # 3) Rename the sheet tab last (so we don't lose the reference)
        ws.title = new_name
        summary.append((old_name, new_name, sheet_code, renumbered))

    wb.save(WB_PATH)

    print("\n=== Realignment complete ===")
    for old, new, code, count in summary:
        print(f"  • {old}\n     → {new}  ({code}, {count} test cases re-IDed with prefix)")


if __name__ == "__main__":
    main()
