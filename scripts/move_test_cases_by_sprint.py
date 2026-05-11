"""Move new test-case sheets from the master Sprint 1 workbook into the
respective per-sprint workbooks (Sprint2..Sprint5).

Routing rule: a sheet name prefixed with `S<N>-` belongs to Sprint <N>.
Original sheets (AUTH-*, ADMIN-08, STU-02, S1-*) stay in Sprint 1.

For each move:
  1) Drop any sheet with the same name in the destination workbook.
  2) Create a new sheet at the end with column widths, merged cells,
     freeze panes, and per-cell value+styles preserved.
  3) Save destination, then remove the source sheet from Sprint 1.
"""

from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection

ROOT = r"e:\UpLoad\capstoneprojectedvision\Ed_Vision"
MASTER = rf"{ROOT}\C1SE.14-Test-Case-Sprint1.xlsx"


def target_path(sprint: int) -> str:
    return rf"{ROOT}\C1SE.14-Test-Case-Sprint{sprint}.xlsx"


def copy_sheet(src_ws, dst_wb, new_name: str) -> None:
    """Copy `src_ws` into `dst_wb` as a brand-new sheet named `new_name`."""
    if new_name in dst_wb.sheetnames:
        del dst_wb[new_name]
    dst_ws = dst_wb.create_sheet(new_name)

    # Column widths
    for col_letter, dim in src_ws.column_dimensions.items():
        if dim.width is not None:
            dst_ws.column_dimensions[col_letter].width = dim.width

    # Row heights + cells with styles
    max_row = src_ws.max_row
    max_col = src_ws.max_column
    for row in src_ws.iter_rows(min_row=1, max_row=max_row,
                                min_col=1, max_col=max_col):
        for cell in row:
            new_cell = dst_ws.cell(row=cell.row, column=cell.column,
                                   value=cell.value)
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.fill = copy(cell.fill)
                new_cell.border = copy(cell.border)
                new_cell.alignment = copy(cell.alignment)
                new_cell.number_format = cell.number_format
                new_cell.protection = copy(cell.protection)

    for r, dim in src_ws.row_dimensions.items():
        if dim.height is not None:
            dst_ws.row_dimensions[r].height = dim.height

    # Merged ranges
    for rng in list(src_ws.merged_cells.ranges):
        dst_ws.merge_cells(str(rng))

    # Freeze panes
    if src_ws.freeze_panes:
        dst_ws.freeze_panes = src_ws.freeze_panes


def main() -> None:
    wb_master = load_workbook(MASTER)

    # Decide routing for each sheet in master
    moves: dict[int, list[str]] = {2: [], 3: [], 4: [], 5: []}
    for name in list(wb_master.sheetnames):
        if not name.startswith("S"):
            continue
        # Names look like "S2-..." / "S3-..." / "S4-..." / "S5-..." / "S1-..."
        if len(name) < 2 or not name[1].isdigit():
            continue
        sprint = int(name[1])
        if sprint == 1:
            continue  # stays in Sprint1
        if sprint in moves:
            moves[sprint].append(name)

    moved_summary = []

    # Per-sprint copy
    for sprint, sheet_names in moves.items():
        if not sheet_names:
            continue
        dst_path = target_path(sprint)
        wb_dst = load_workbook(dst_path)
        for sn in sheet_names:
            src_ws = wb_master[sn]
            copy_sheet(src_ws, wb_dst, sn)
            moved_summary.append((sprint, sn))
        wb_dst.save(dst_path)
        print(f"  → Sprint {sprint}: copied {len(sheet_names)} sheets")

    # Now delete moved sheets from master
    for _, sn in moved_summary:
        if sn in wb_master.sheetnames:
            del wb_master[sn]
    wb_master.save(MASTER)

    print("\n=== Move complete ===")
    by_sprint: dict[int, list[str]] = {}
    for sprint, sn in moved_summary:
        by_sprint.setdefault(sprint, []).append(sn)
    for sprint in sorted(by_sprint):
        print(f"\nSprint {sprint} (+{len(by_sprint[sprint])} sheets):")
        for sn in by_sprint[sprint]:
            print(f"   • {sn}")

    print("\nSprint 1 retained sheets:")
    for sn in wb_master.sheetnames:
        print(f"   • {sn}")


if __name__ == "__main__":
    main()
