from docx import Document
d = Document(r'e:\UpLoad\capstoneprojectedvision\Ed_Vision\C2SE.76_TestPlan_Ver1.1.docx')
out = []
# Walk top-level body to keep paragraph & table order
from docx.oxml.ns import qn
body = d.element.body
for child in body.iterchildren():
    if child.tag == qn('w:p'):
        # find matching paragraph
        for p in d.paragraphs:
            if p._p is child:
                txt = p.text.strip()
                if txt:
                    out.append('P: ' + txt)
                break
    elif child.tag == qn('w:tbl'):
        for t in d.tables:
            if t._tbl is child:
                out.append('--TABLE--')
                for r in t.rows:
                    out.append(' | '.join(c.text.strip().replace('\n',' / ') for c in r.cells))
                out.append('--ENDTABLE--')
                break
print('\n'.join(out))
