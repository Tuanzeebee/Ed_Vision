from openpyxl import load_workbook
wb = load_workbook(r'e:\UpLoad\capstoneprojectedvision\Ed_Vision\C2SE.76_Product&SprintBacklog.xlsx', data_only=True)
for sn in ['Product Backlog', 'Sprint 1 Backlog', 'Sprint 2 Backlog', 'Sprint 3 Backlog', 'Sprint 4 Backlog', 'Sprint 5 Backlog']:
    ws = wb[sn]
    print('=' * 80)
    print('SHEET:', sn, 'dims=', ws.dimensions)
    for r in ws.iter_rows(values_only=True):
        # filter out totally empty
        if any(c not in (None, '') for c in r):
            print(' | '.join('' if c is None else str(c)[:60] for c in r))
